"""Builds the Stage-Pay tracker workbook (blank template and filled example)."""
import sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter

BRICK, INK, MUTED, LINE = "B4492B", "1D2321", "5B6360", "D9D8D1"
INPUT_FILL = PatternFill("solid", fgColor="FFF4CC")
HEAD_FILL = PatternFill("solid", fgColor=BRICK)
SOFT_FILL = PatternFill("solid", fgColor="F3F2EE")
OK_FILL, OK_FONT = PatternFill("solid", fgColor="E1EEE6"), Font(name="Arial", color="2F6B4F", bold=True)
WAIT_FILL, WAIT_FONT = PatternFill("solid", fgColor="F7ECD2"), Font(name="Arial", color="9A6A0B", bold=True)
READY_FILL, READY_FONT = PatternFill("solid", fgColor="F6E3DC"), Font(name="Arial", color=BRICK, bold=True)
LOCK_FONT = Font(name="Arial", color="7B8380")
THIN = Side(style="thin", color=LINE)
BOX = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)
USD = '$#,##0;($#,##0);"-"'
WRAP = Alignment(wrap_text=True, vertical="top")

def F(size=10, bold=False, color=INK, italic=False):
    return Font(name="Arial", size=size, bold=bold, color=color, italic=italic)

def title(ws, text, sub=None):
    ws["A1"] = text; ws["A1"].font = F(18, True)
    if sub: ws["A2"] = sub; ws["A2"].font = F(10, color=MUTED)
    ws.sheet_view.showGridLines = False

def header(ws, row, labels, col=1):
    for i, h in enumerate(labels):
        c = ws.cell(row=row, column=col + i, value=h)
        c.font = F(10, True, "FFFFFF"); c.fill = HEAD_FILL; c.alignment = Alignment(vertical="center", wrap_text=True); c.border = BOX
    ws.row_dimensions[row].height = 30

def widths(ws, ws_widths):
    for i, w in enumerate(ws_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

def inp(c, fmt=None):
    c.fill = INPUT_FILL; c.border = BOX; c.font = F(10, color="1F3FBF")
    if fmt: c.number_format = fmt

def calc(c, fmt=None, bold=False):
    c.border = BOX; c.font = F(10, bold)
    if fmt: c.number_format = fmt

STAGES = [
    ("Foundation & footings", [
        "Setting-out photo: pegs and strings match the approved plan",
        "Trenches with a tape measure showing depth and width",
        "Footings poured: photo of every trench",
        "Foundation brickwork up to ground level: photo of each side",
        "Video walking round the whole foundation",
        "Receipts: cement, sand, stone, bricks"]),
    ("Slab level (DPC & floor slab)", [
        "Brickwork to slab level: photo of each side",
        "Damp-proof course (DPC) visible all round",
        "Hardcore filled and compacted",
        "Slab poured: photo across the whole floor",
        "Receipts: cement, mesh, DPC, hardcore",
        "Council stage inspection signed (if your council requires it)"]),
    ("Window level", [
        "Photo of each outside wall",
        "Window frames set in place",
        "Door frames set in place",
        "Video walking through every room",
        "Receipts: bricks, cement, frames"]),
    ("Roof level (lintels & wall plate)", [
        "Lintels over every door and window",
        "Ring beam / brickforce at wall plate",
        "Walls at full height: photo of each side",
        "Video walk-round inside and outside",
        "Receipts: steel, cement, bricks"]),
    ("Roofing", [
        "Trusses up: photo from two corners",
        "Roof covering complete: photo from the road",
        "Fascia boards and gutters fitted",
        "Photo of the roof from inside, before ceilings",
        "Receipts: timber, sheets/tiles, nails, straps"]),
    ("Plumbing & electrical (first fix)", [
        "Conduits and boxes in the walls: photo of each room",
        "Water pipes and drainage laid",
        "Septic tank or sewer connection: photo",
        "Name and phone number of the electrician and plumber",
        "Receipts: cables, pipes, fittings"]),
    ("Plaster & ceilings", [
        "Inside plaster done: video of every room",
        "Outside plaster done: photo of each side",
        "Ceilings fitted: photo of each room",
        "Receipts: cement, plaster sand, ceiling boards"]),
    ("Floors, doors & finishing", [
        "Floor screed or tiles done: video of every room",
        "Doors hung and locks fitted",
        "Window glass fitted",
        "Painting done inside and outside",
        "Sinks, toilet, lights and switches installed",
        "Final receipts and handover of keys"]),
]
N_STAGE_ROWS = 10           # 8 standard stages + 2 spare rows
S0 = 6                      # first data row on Stages
SL = S0 + N_STAGE_ROWS - 1  # last stage row
P0, PL = 6, 205             # Payments rows
C0, CL = 6, 155             # Proof Checklist rows

EXAMPLE = {
    "project": ["4-room house", "Stand 1234, Chitungwiza", "Tendai M.", "+263 77 000 0000",
                "Brother: Farai (visits on Saturdays)", "2026-03-01", "USD"],
    "agreed": [5400, 6200, 4800, 3000, 7500, 2500, 3500, 6000],
    "received": {1: 6, 2: 6, 3: 2},   # stage -> number of proof items ticked (in order)
    "payments": [
        ("2026-03-04", 1, 3000, "Mukuru", "Tendai M.", "Foundation materials"),
        ("2026-03-21", 1, 2400, "Bank transfer", "Tendai M.", "Foundation labour"),
        ("2026-04-15", 2, 4000, "Mukuru", "Tendai M.", "Bricks and cement for slab level"),
        ("2026-05-02", 2, 2200, "Bank transfer", "Tendai M.", "Slab and labour"),
        ("2026-06-10", 3, 2600, "WorldRemit", "Tendai M.", "Bricks to window level"),
    ],
}

def build(path, example):
    import datetime as dt
    wb = Workbook()

    # ---------------- Start Here ----------------
    ws = wb.active; ws.title = "Start Here"
    title(ws, "Stage-Pay Tracker", "Build at home from abroad. No proof, no payment.")
    widths(ws, [4, 30, 80])
    rows = [
        ("h", "How it works"),
        ("1", "Fill in the Project tab: your builder's name, location and start date."),
        ("2", "On the Stages tab, enter the agreed price for each stage. Rename or adjust stages to match your contract."),
        ("3", "Check the Proof Checklist. Each stage has a list of photos, videos and receipts to ask for. Add or remove items to suit your build."),
        ("4", "Before paying for a stage, open the WhatsApp tab. It writes the message asking for any missing proof. Copy it and send it to your builder."),
        ("5", "As proof arrives, set 'Received?' to Yes on the Proof Checklist. Paste a link to where you saved the photo or video."),
        ("6", "Every time you send money, add a row to the Payments tab. The Dashboard updates by itself."),
        ("h", "The rule"),
        ("•", "A stage is Verified only when money has been sent AND every proof item is received."),
        ("•", "The next stage stays Locked until the previous one is Verified. Don't send money to a locked stage."),
        ("•", "'Unproven' on the Dashboard is money you've sent that isn't backed by proof yet. Aim to keep it at $0 before each new payment."),
        ("h", "Colour key"),
        ("", "Yellow cells: type here"),
        ("", "White cells: worked out for you, so don't type over them"),
        ("h", "Tabs"),
        ("", "Dashboard: totals, progress and your next step"),
        ("", "Project: your project details"),
        ("", "Stages: agreed price, money sent and status for each stage"),
        ("", "Payments: every transfer you make"),
        ("", "Proof Checklist: what to ask for at each stage, and what you've received"),
        ("", "WhatsApp: ready-to-send messages for your builder and family"),
        ("", "Before You Build: checks to do first, and red flags to watch for"),
        ("h", "Using Google Sheets"),
        ("", "Upload this file to Google Drive, then open it with Google Sheets (or go to File > Import in a new sheet). Everything works on a phone in the Google Sheets app."),
        ("h", "Important"),
        ("", "Stage-Pay is a record-keeping tool. It isn't legal, financial or engineering advice, and it can't guarantee how a builder behaves. For large projects, consider an independent inspector or engineer to check the work on site."),
    ]
    r = 4
    for a, b in rows:
        if a == "h":
            r += 1; ws.cell(row=r, column=2, value=b).font = F(12, True, BRICK)
        else:
            ws.cell(row=r, column=2 if a == "" else 1, value=a if a else b).font = F(10, a in ("1","2","3","4","5","6"), BRICK if a else INK)
            c = ws.cell(row=r, column=3 if a else 2, value=b if a else None)
            if a: c.font = F(10); c.alignment = WRAP
            else:
                ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
                ws.cell(row=r, column=2).alignment = WRAP; ws.cell(row=r, column=2).font = F(10)
        r += 1
    # colour key swatches
    for rr in range(4, r):
        v = ws.cell(row=rr, column=2).value
        if v == "Yellow cells: type here": ws.cell(row=rr, column=2).fill = INPUT_FILL
        if v and v.startswith("White cells"): ws.cell(row=rr, column=2).border = BOX
    for rr in range(4, r):
        v = ws.cell(row=rr, column=2).value or ws.cell(row=rr, column=3).value or ""
        if len(str(v)) > 90: ws.row_dimensions[rr].height = 28 if len(str(v)) < 170 else 42

    # ---------------- Project ----------------
    pj = wb.create_sheet("Project")
    title(pj, "Project details", "Fill in the yellow cells.")
    widths(pj, [28, 46, 50])
    labels = ["Project name", "Location / stand number", "Builder name", "Builder WhatsApp number",
              "Site contact (family or friend)", "Start date", "Currency"]
    hints = ["e.g. 4-room house", "e.g. Stand 1234, Chitungwiza", "Used in your WhatsApp messages",
             "", "Someone you trust who can visit the site", "", "All amounts in this tracker use this currency"]
    for i, (l, h) in enumerate(zip(labels, hints)):
        rr = 4 + i
        pj.cell(row=rr, column=1, value=l).font = F(10, True)
        c = pj.cell(row=rr, column=2); inp(c)
        if example:
            v = EXAMPLE["project"][i]
            c.value = dt.date.fromisoformat(v) if l == "Start date" else v
        elif l == "Currency": c.value = "USD"
        if l == "Start date": c.number_format = "d mmm yyyy"
        pj.cell(row=rr, column=3, value=h).font = F(9, color=MUTED, italic=True)
    BUILDER = 'IF(Project!$B$6="","there",Project!$B$6)'

    # ---------------- Stages ----------------
    st = wb.create_sheet("Stages")
    title(st, "Stages", "Enter the agreed price for each stage (yellow). Everything else updates from the Payments and Proof Checklist tabs.")
    header(st, 5, ["#", "Stage", "Agreed price", "Sent", "Left to pay", "Proof items", "Proof received", "Proof %", "Status", "Next step", "Open?"])
    widths(st, [5, 36, 14, 14, 14, 11, 12, 10, 16, 58, 7])
    for i in range(N_STAGE_ROWS):
        r = S0 + i
        st.cell(row=r, column=1, value=i + 1).font = F(10, True)
        st.cell(row=r, column=1).border = BOX
        name = st.cell(row=r, column=2, value=STAGES[i][0] if i < len(STAGES) else None); inp(name)
        ag = st.cell(row=r, column=3); inp(ag, USD)
        if example and i < len(EXAMPLE["agreed"]): ag.value = EXAMPLE["agreed"][i]
        calc(st.cell(row=r, column=4, value=f"=SUMIF(Payments!$B${P0}:$B${PL},A{r},Payments!$D${P0}:$D${PL})"), USD)
        calc(st.cell(row=r, column=5, value=f'=IF(B{r}="","",C{r}-D{r})'), USD)
        calc(st.cell(row=r, column=6, value=f"=COUNTIFS('Proof Checklist'!$A${C0}:$A${CL},A{r},'Proof Checklist'!$C${C0}:$C${CL},\"<>\")"))
        calc(st.cell(row=r, column=7, value=f"=COUNTIFS('Proof Checklist'!$A${C0}:$A${CL},A{r},'Proof Checklist'!$D${C0}:$D${CL},\"Yes\")"))
        calc(st.cell(row=r, column=8, value=f'=IF(F{r}=0,0,G{r}/F{r})'), "0%")
        prev_ok = "TRUE" if i == 0 else f'I{r-1}="Verified"'
        calc(st.cell(row=r, column=9, value=(
            f'=IF(B{r}="","",IF(F{r}=0,"Add proof items",IF(AND(D{r}>0,G{r}>=F{r}),"Verified",'
            f'IF(NOT({prev_ok}),"Locked",IF(D{r}>0,"Hold payment","Ready to fund")))))')), bold=True)
        calc(st.cell(row=r, column=10, value=(
            f'=IF(I{r}="","",IF(I{r}="Verified","Done. All proof received.",IF(I{r}="Locked","Wait. Don\'t send money until the previous stage is verified.",'
            f'IF(I{r}="Hold payment","Hold. "&(F{r}-G{r})&" proof item(s) missing. Send the WhatsApp request.",'
            f'IF(I{r}="Ready to fund","Agree the amount and send proof list before paying.","Add proof items on the Proof Checklist tab.")))))')))
        st.cell(row=r, column=10).alignment = Alignment(wrap_text=True, vertical="center")
        calc(st.cell(row=r, column=11, value=f'=IF(AND(B{r}<>"",I{r}<>"Verified"),1,0)'))
        st.cell(row=r, column=11).font = F(9, color=MUTED)
    tr = SL + 1
    st.cell(row=tr, column=2, value="Total").font = F(10, True)
    for col in (3, 4, 5):
        L = get_column_letter(col)
        calc(st.cell(row=tr, column=col, value=f"=SUM({L}{S0}:{L}{SL})"), USD, bold=True)
    st.cell(row=tr + 2, column=2, value="Rows 9 and 10 are spare. Type a stage name to use them. 'Open?' is a helper column the Dashboard uses.").font = F(9, color=MUTED, italic=True)
    st.freeze_panes = "C6"
    rng = f"I{S0}:I{SL}"
    st.conditional_formatting.add(rng, FormulaRule(formula=[f'$I{S0}="Verified"'], fill=OK_FILL, font=OK_FONT))
    st.conditional_formatting.add(rng, FormulaRule(formula=[f'$I{S0}="Hold payment"'], fill=WAIT_FILL, font=WAIT_FONT))
    st.conditional_formatting.add(rng, FormulaRule(formula=[f'$I{S0}="Ready to fund"'], fill=READY_FILL, font=READY_FONT))
    st.conditional_formatting.add(rng, FormulaRule(formula=[f'$I{S0}="Locked"'], font=LOCK_FONT))
    STAGE_NAME = f"Stages!$B${S0}:$B${SL}"
    STAGE_STATUS = f"Stages!$I${S0}:$I${SL}"

    # ---------------- Payments ----------------
    pm = wb.create_sheet("Payments")
    title(pm, "Payments", "Add a row every time you send money. Pick the stage number (1–10) the money is for.")
    header(pm, 5, ["Date", "Stage #", "Stage", "Amount", "Method", "Paid to", "What it was for / reference", "Check"])
    widths(pm, [13, 9, 32, 13, 17, 20, 38, 34])
    dv_stage = DataValidation(type="whole", operator="between", formula1="1", formula2=str(N_STAGE_ROWS), allow_blank=True,
                              showErrorMessage=True, errorTitle="Stage number", error="Enter a stage number from 1 to 10.")
    dv_method = DataValidation(type="list", formula1='"Mukuru,WorldRemit,Bank transfer,EcoCash,InnBucks,Cash via family,Paid supplier directly,Other"', allow_blank=True)
    pm.add_data_validation(dv_stage); pm.add_data_validation(dv_method)
    for r in range(P0, PL + 1):
        for col, fmt in ((1, "d mmm yyyy"), (2, "0"), (4, USD), (5, None), (6, None), (7, None)):
            inp(pm.cell(row=r, column=col), fmt)
        calc(pm.cell(row=r, column=3, value=f'=IF(B{r}="","",IFERROR(INDEX({STAGE_NAME},B{r}),"Unknown stage"))'))
        calc(pm.cell(row=r, column=8, value=(
            f'=IF(B{r}="","",IF(B{r}=1,"OK",IF(IFERROR(INDEX({STAGE_STATUS},B{r}-1),"")="Verified","OK","⚠ Previous stage not verified yet")))')))
    dv_stage.add(f"B{P0}:B{PL}"); dv_method.add(f"E{P0}:E{PL}")
    pm.conditional_formatting.add(f"H{P0}:H{PL}", FormulaRule(formula=[f'LEFT($H{P0},1)="⚠"'], fill=WAIT_FILL, font=WAIT_FONT))
    pm.conditional_formatting.add(f"H{P0}:H{PL}", FormulaRule(formula=[f'$H{P0}="OK"'], font=OK_FONT))
    pm["J5"] = "Total sent"; pm["J5"].font = F(10, True)
    calc(pm.cell(row=6, column=10, value=f"=SUM(D{P0}:D{PL})"), USD, bold=True)
    pm.column_dimensions["I"].width = 3; pm.column_dimensions["J"].width = 14
    pm.freeze_panes = "A6"
    if example:
        for k, (d, s, a, m, to, note) in enumerate(EXAMPLE["payments"]):
            r = P0 + k
            pm.cell(row=r, column=1, value=dt.date.fromisoformat(d)); pm.cell(row=r, column=2, value=s)
            pm.cell(row=r, column=4, value=a); pm.cell(row=r, column=5, value=m)
            pm.cell(row=r, column=6, value=to); pm.cell(row=r, column=7, value=note)
    else:
        pm["A4"] = "Example of a filled row:  4 Mar 2026  |  1  |  (stage fills in)  |  $3,000  |  Mukuru  |  Tendai M.  |  Foundation materials"
        pm["A4"].font = F(9, color=MUTED, italic=True)

    # ---------------- Proof Checklist ----------------
    pc = wb.create_sheet("Proof Checklist")
    title(pc, "Proof Checklist", "What to ask for before each stage is paid. Set 'Received?' to Yes when it arrives. Add your own items in the empty rows at the bottom.")
    header(pc, 5, ["Stage #", "Stage", "Proof to ask for", "Received?", "Date received", "Where it's saved (link or note)", "msg"])
    widths(pc, [9, 30, 58, 11, 14, 40, 5])
    dv_yes = DataValidation(type="list", formula1='"Yes,No"', allow_blank=True)
    dv_st2 = DataValidation(type="whole", operator="between", formula1="1", formula2=str(N_STAGE_ROWS), allow_blank=True)
    pc.add_data_validation(dv_yes); pc.add_data_validation(dv_st2)
    r = C0
    for si, (_, items) in enumerate(STAGES, 1):
        got = EXAMPLE["received"].get(si, 0) if example else 0
        for k, item in enumerate(items):
            pc.cell(row=r, column=1, value=si); pc.cell(row=r, column=3, value=item)
            pc.cell(row=r, column=4, value="Yes" if k < got else "No")
            if example and k < got:
                pc.cell(row=r, column=5, value=dt.date(2026, 3 + si, 10 + k))
                pc.cell(row=r, column=6, value=f"Google Drive > Build > Stage {si}")
            r += 1
    for r in range(C0, CL + 1):
        inp(pc.cell(row=r, column=1), "0"); inp(pc.cell(row=r, column=3)); inp(pc.cell(row=r, column=4))
        inp(pc.cell(row=r, column=5), "d mmm yyyy"); inp(pc.cell(row=r, column=6))
        pc.cell(row=r, column=3).alignment = Alignment(wrap_text=True, vertical="center")
        calc(pc.cell(row=r, column=2, value=f'=IF(A{r}="","",IFERROR(INDEX({STAGE_NAME},A{r}),""))'))
        # helper: numbers the missing items for the stage selected on the WhatsApp tab
        g = pc.cell(row=r, column=7, value=(
            f'=IF(AND(A{r}=WhatsApp!$C$5,C{r}<>"",D{r}<>"Yes"),COUNTIFS($A${C0}:A{r},WhatsApp!$C$5,$C${C0}:C{r},"<>",$D${C0}:D{r},"<>Yes"),"")'))
        g.font = F(8, color=MUTED)
    dv_yes.add(f"D{C0}:D{CL}"); dv_st2.add(f"A{C0}:A{CL}")
    pc.conditional_formatting.add(f"D{C0}:D{CL}", FormulaRule(formula=[f'$D{C0}="Yes"'], fill=OK_FILL, font=OK_FONT))
    pc.conditional_formatting.add(f"A{C0}:F{CL}", FormulaRule(formula=[f'AND($A{C0}<>"",ISODD($A{C0}))'], fill=SOFT_FILL))
    pc.freeze_panes = "A6"

    # ---------------- WhatsApp ----------------
    wa = wb.create_sheet("WhatsApp")
    title(wa, "WhatsApp messages", "Copy a message, paste it into WhatsApp and send. Your builder's name comes from the Project tab.")
    widths(wa, [4, 30, 95])
    wa["B5"] = "Stage to ask about"; wa["B5"].font = F(10, True)
    c = wa["C5"]; c.value = f"=IFERROR(MATCH(1,Stages!$K${S0}:$K${SL},0),{N_STAGE_ROWS})"; inp(c, "0")
    c.alignment = Alignment(horizontal="left")
    wa["D5"] = "← Picks your current stage. Type a different number (1–10) to change it."
    wa["D5"].font = F(9, color=MUTED, italic=True)
    wa["B6"] = "Stage name"; wa["B6"].font = F(10, True)
    calc(wa.cell(row=6, column=3, value=f'=IFERROR(INDEX({STAGE_NAME},C5),"")'), bold=True)
    wa["B8"] = "Missing proof"; wa["B8"].font = F(12, True, BRICK)
    for k in range(1, 11):
        r = 8 + k
        wa.cell(row=r, column=2, value=k).font = F(9, color=MUTED)
        calc(wa.cell(row=r, column=3, value=(
            f"=IFERROR(INDEX('Proof Checklist'!$C${C0}:$C${CL},MATCH({k},'Proof Checklist'!$G${C0}:$G${CL},0)),\"\")")))
    wa["B20"] = "Message: ask for missing proof"; wa["B20"].font = F(12, True, BRICK)
    lines = "&".join([f'IF(C{8+k}="","",CHAR(10)&"{k}. "&C{8+k})' for k in range(1, 11)])
    msg = wa.cell(row=21, column=3, value=(
        f'=IF(C9="","Hi "&{BUILDER}&", thank you. All the proof for "&LOWER(C6)&" is in. I\'ll send the payment today.",'
        f'"Hi "&{BUILDER}&", hope you and the team are well. Before I send the next payment for "&LOWER(C6)&", please send:"&{lines}&CHAR(10)&CHAR(10)&"Thank you!")'))
    msg.alignment = WRAP; msg.border = BOX; msg.font = F(10)
    wa.row_dimensions[21].height = 190
    wa["B21"] = "Copy the cell on the right →"; wa["B21"].font = F(9, color=MUTED, italic=True); wa["B21"].alignment = WRAP

    templates = [
        ("Before work starts: agree the rule",
         '="Hi "&{b}&", before we start, here\'s how I\'ll handle payments so we\'re always clear. For each stage I\'ll send you a short list of photos, a video and receipts. Once I have them, I send the money for the next stage the same day. It\'s the same for every stage and everyone I work with. Is that OK with you?"'),
        ("Proof is late",
         '="Hi "&{b}&", just checking in on the photos and video for "&LOWER(C6)&". I\'m ready to send the next payment as soon as they come through."'),
        ("Holding a payment politely",
         '="Hi "&{b}&", thanks for the update. I\'m still missing a few items for "&LOWER(C6)&" (see my list above). I\'ll release the payment as soon as they\'re in, so we keep to the agreed system."'),
        ("Price has gone up",
         '="Hi "&{b}&", understood. Please send the supplier\'s quote or receipt showing the new price. Once I see it, I\'ll update the budget and we can agree the new amount."'),
        ("For family managing the site",
         '="Hi, thank you so much for keeping an eye on the build, it means a lot. To keep things simple and fair for everyone, I\'m using the same system for every stage: photos, a short video and receipts before each payment. Could you help me get these for "&LOWER(C6)&"?"'),
        ("Payment sent",
         '="Hi "&{b}&", I\'ve sent the payment for "&LOWER(C6)&" today. Please confirm when you receive it. Thank you for the good work!"'),
    ]
    r = 24
    wa.cell(row=r - 1, column=2, value="More ready-made messages").font = F(12, True, BRICK)
    for name, f in templates:
        wa.cell(row=r, column=2, value=name).font = F(10, True); wa.cell(row=r, column=2).alignment = WRAP
        c = wa.cell(row=r, column=3, value=f.format(b=BUILDER)); c.alignment = WRAP; c.border = BOX; c.font = F(10)
        wa.row_dimensions[r].height = 62
        r += 1

    # ---------------- Before You Build ----------------
    bb = wb.create_sheet("Before You Build")
    title(bb, "Before you build", "Do these checks before you send the first payment. Mark each one Yes when it's done.")
    widths(bb, [4, 70, 11, 44])
    header(bb, 5, ["Check", "Done?", "Notes"], col=2)
    checks = [
        "The stand is legally yours: title deed, or cession / agreement of sale, checked with the Deeds Registry, the council or the developer, in YOUR name",
        "You've seen the stand's beacons (in person, by someone you trust, or on a live video call)",
        "You know what servicing is done (water, sewer, roads) and what it will cost if it isn't",
        "Building plans are approved by the local council before any work starts",
        "You've met the builder (in person or on video) and seen at least 2 houses they've built",
        "You've spoken to at least one past client of the builder",
        "Written agreement signed: what will be built, price per stage, timeline, who buys materials, what happens if work stops",
        "Payment schedule agreed: money per stage, and the next stage paid only after proof",
        "Proof list sent to the builder before work starts (WhatsApp tab, 'Before work starts')",
        "Someone you trust, who isn't the builder, can visit the site now and then",
        "Where possible, you pay material suppliers directly and keep their receipts",
        "One folder (e.g. Google Drive) for all photos, videos, receipts and agreements",
    ]
    for i, t in enumerate(checks):
        r = 6 + i
        c = bb.cell(row=r, column=2, value=t); c.alignment = WRAP; c.border = BOX; c.font = F(10)
        d = bb.cell(row=r, column=3, value="No"); inp(d)
        inp(bb.cell(row=r, column=4)); bb.cell(row=r, column=4).alignment = WRAP
        bb.row_dimensions[r].height = 30 if len(t) > 70 else 18
    dv_b = DataValidation(type="list", formula1='"Yes,No"', allow_blank=True); bb.add_data_validation(dv_b)
    dv_b.add(f"C6:C{5+len(checks)}")
    bb.conditional_formatting.add(f"C6:C{5+len(checks)}", FormulaRule(formula=['$C6="Yes"'], fill=OK_FILL, font=OK_FONT))
    r = 6 + len(checks) + 1
    bb.cell(row=r, column=2, value="Checks done").font = F(10, True)
    calc(bb.cell(row=r, column=3, value=f'=COUNTIF(C6:C{5+len(checks)},"Yes")&" of {len(checks)}"'), bold=True)
    r += 2
    bb.cell(row=r, column=2, value="Red flags: stop and ask questions if you see these").font = F(12, True, BRICK)
    flags = [
        "Asks for most of the money up front, before any work",
        "Asks you to send money to a different person or account than you agreed",
        "Photos are always close-ups, from the same angle, or there's never a video",
        "Says 'prices went up' but can't show a supplier quote or receipt",
        "Work has stalled, but requests for money keep coming",
        "Pressure to pay today: 'the materials will be gone' or 'the price changes tomorrow'",
        "Becomes hard to reach right after you send a payment",
        "Seller can't show the stand beacons, or rushes you to pay before you've checked ownership",
    ]
    for t in flags:
        r += 1
        bb.cell(row=r, column=2, value="⚑  " + t).font = F(10); bb.cell(row=r, column=2).alignment = WRAP
    bb.freeze_panes = "A6"

    # ---------------- Dashboard ----------------
    db = wb.create_sheet("Dashboard", 1)
    db.sheet_view.showGridLines = False
    widths(db, [3, 22, 18, 18, 18, 18, 3])
    db["B1"] = "=IF(Project!$B$4=\"\",\"My build\",Project!$B$4)"; db["B1"].font = F(18, True)
    db["B2"] = '=IF(Project!$B$5="","Fill in the Project tab to get started",Project!$B$5&"  ·  Builder: "&Project!$B$6)'
    db["B2"].font = F(10, color=MUTED)
    db["B4"], db["C4"], db["D4"], db["E4"], db["F4"] = "Total budget", "Sent", "Verified built", "Unproven", "Budget left"
    db["B5"] = f"=SUM(Stages!C{S0}:C{SL})"
    db["C5"] = f"=SUM(Stages!D{S0}:D{SL})"
    db["D5"] = f'=SUMIF(Stages!I{S0}:I{SL},"Verified",Stages!D{S0}:D{SL})'
    db["E5"] = "=C5-D5"
    db["F5"] = "=B5-C5"
    for col in range(2, 7):
        l = db.cell(row=4, column=col); l.font = F(9, True, MUTED); l.fill = SOFT_FILL; l.border = BOX
        v = db.cell(row=5, column=col); v.font = F(16, True); v.number_format = USD; v.border = BOX
    db.row_dimensions[5].height = 30
    db.conditional_formatting.add("E5", FormulaRule(formula=["$E$5>0"], font=Font(name="Arial", size=16, bold=True, color=BRICK)))
    db.conditional_formatting.add("E5", FormulaRule(formula=["$E$5<=0"], font=Font(name="Arial", size=16, bold=True, color="2F6B4F")))

    db["B7"] = "Verified"; db["B8"] = "Sent, no proof yet"
    db["C7"] = '=IF($B$5=0,0,$D$5/$B$5)'; db["C8"] = '=IF($B$5=0,0,$E$5/$B$5)'
    db["D7"] = '=REPT("█",ROUND(C7*40,0))'; db["D8"] = '=REPT("█",ROUND(C8*40,0))'
    for rr, color in ((7, "2F6B4F"), (8, BRICK)):
        db.cell(row=rr, column=2).font = F(10, True)
        db.cell(row=rr, column=3).number_format = "0%"; db.cell(row=rr, column=3).font = F(10, True)
        db.cell(row=rr, column=4).font = Font(name="Arial", size=10, color=color)
    db["B9"] = "(Bars show share of total budget)"; db["B9"].font = F(8, color=MUTED, italic=True)

    db["B11"] = "Where you are now"; db["B11"].font = F(12, True, BRICK)
    db["B12"] = "Current stage"; db["B13"] = "Status"; db["B14"] = "Next step"; db["B15"] = "Payment warnings"
    db["C12"] = f'=IFERROR(INDEX(Stages!B{S0}:B{SL},MATCH(1,Stages!K{S0}:K{SL},0)),"All stages verified 🎉")'
    db["C13"] = f'=IFERROR(INDEX(Stages!I{S0}:I{SL},MATCH(1,Stages!K{S0}:K{SL},0)),"Verified")'
    db["C14"] = f'=IFERROR(INDEX(Stages!J{S0}:J{SL},MATCH(1,Stages!K{S0}:K{SL},0)),"Well done. Keep all receipts and photos safe.")'
    db["C15"] = f'=IF(COUNTIF(Payments!H{P0}:H{PL},"⚠*")=0,"None","⚠ "&COUNTIF(Payments!H{P0}:H{PL},"⚠*")&" payment(s) sent before the previous stage was verified")'
    for rr in range(12, 16):
        db.cell(row=rr, column=2).font = F(10, True)
        db.merge_cells(start_row=rr, start_column=3, end_row=rr, end_column=6)
        db.cell(row=rr, column=3).font = F(10, rr == 12); db.cell(row=rr, column=3).alignment = Alignment(wrap_text=True, vertical="center")
    db.row_dimensions[14].height = 30
    db.conditional_formatting.add("C13", FormulaRule(formula=['$C$13="Hold payment"'], fill=WAIT_FILL, font=WAIT_FONT))
    db.conditional_formatting.add("C13", FormulaRule(formula=['$C$13="Ready to fund"'], fill=READY_FILL, font=READY_FONT))
    db.conditional_formatting.add("C13", FormulaRule(formula=['$C$13="Verified"'], fill=OK_FILL, font=OK_FONT))
    db.conditional_formatting.add("C15", FormulaRule(formula=['LEFT($C$15,1)="⚠"'], fill=WAIT_FILL, font=WAIT_FONT))

    db["B17"] = "All stages"; db["B17"].font = F(12, True, BRICK)
    header(db, 18, ["Stage", "Agreed", "Sent", "Proof", "Status"], col=2)
    for i in range(N_STAGE_ROWS):
        r, s = 19 + i, S0 + i
        db.cell(row=r, column=2, value=f'=IF(Stages!B{s}="","",Stages!A{s}&". "&Stages!B{s})')
        db.cell(row=r, column=3, value=f'=IF(Stages!B{s}="","",Stages!C{s})').number_format = USD
        db.cell(row=r, column=4, value=f'=IF(Stages!B{s}="","",Stages!D{s})').number_format = USD
        db.cell(row=r, column=5, value=f'=IF(Stages!B{s}="","",Stages!G{s}&" of "&Stages!F{s})')
        db.cell(row=r, column=6, value=f"=Stages!I{s}")
        for col in range(2, 7):
            c = db.cell(row=r, column=col); c.border = BOX; c.font = F(10, col == 6)
            if col == 2: c.alignment = Alignment(wrap_text=True, vertical="center")
        db.row_dimensions[r].height = 28
    rng = "F19:F28"
    db.conditional_formatting.add(rng, FormulaRule(formula=['$F19="Verified"'], fill=OK_FILL, font=OK_FONT))
    db.conditional_formatting.add(rng, FormulaRule(formula=['$F19="Hold payment"'], fill=WAIT_FILL, font=WAIT_FONT))
    db.conditional_formatting.add(rng, FormulaRule(formula=['$F19="Ready to fund"'], fill=READY_FILL, font=READY_FONT))
    db.conditional_formatting.add(rng, FormulaRule(formula=['$F19="Locked"'], font=LOCK_FONT))
    db.column_dimensions["B"].width = 34

    # tab colours
    for name, col in (("Start Here", INK), ("Dashboard", BRICK), ("Project", "C9A227"), ("Stages", "C9A227"),
                      ("Payments", "C9A227"), ("Proof Checklist", "C9A227"), ("WhatsApp", "2F6B4F"), ("Before You Build", "2F6B4F")):
        wb[name].sheet_properties.tabColor = col
    wb.active = 1 if example else 0
    wb.save(path)

if __name__ == "__main__":
    out = sys.argv[1] if len(sys.argv) > 1 else "."
    build(f"{out}/Stage-Pay-Tracker.xlsx", example=False)
    build(f"{out}/Stage-Pay-Tracker-EXAMPLE.xlsx", example=True)
    print("built")
